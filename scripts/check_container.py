"""Run inside the Compose service: python < scripts/check_container.py.

Uses synthetic content and mock embeddings; no model server or network is needed.
"""
import os
from pathlib import Path
import subprocess
import tempfile

assert os.getuid() != 0, "Container must run as a non-root user"
assert os.statvfs(".").f_flag & os.ST_RDONLY, "Application filesystem must be read-only"
subprocess.run(["git", "--version"], check=True, timeout=10)

# Exercise React asset delivery and API/session persistence without a model.
import httpx
import re
from utils.runtime import request_state

with httpx.Client(base_url="http://127.0.0.1:8765", headers={"X-DocMind-Client": "react"}) as client:
    assert client.get("/api/health").json() == {"ok": True}
    page = client.get("/")
    page.raise_for_status()
    assert '<div id="root"></div>' in page.text
    assets = re.findall(r'(?:src|href)="(/assets/[^"]+)"', page.text)
    assert assets
    for asset in assets:
        client.get(asset).raise_for_status()
    session = client.post("/api/session")
    session.raise_for_status()
    assert session.json()["settings"]["endpoint"] == os.environ["DOCMIND_OLLAMA_ENDPOINT"]
    client.delete("/api/session").raise_for_status()
print("PASS React assets, API health and writable conversation storage", flush=True)

from docx import Document as WordDocument
from pptx import Presentation
from llama_index.core.embeddings import MockEmbedding
from utils.document_readers import libreoffice_executable
from utils.llama_index import create_index, load_documents
from utils.session_workspace import session_directory, clear_session_workspace

marker = "ORBIT-7429"
state = {}
try:
    directory = session_directory("data", state)
    word = WordDocument()
    word.add_paragraph(f"Launch approval code: {marker}")
    word.save(directory / "sample.docx")
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[1]).shapes.title.text = marker
    slides.save(directory / "sample.pptx")
    # Real legacy containers test both Writer/Impress and upload conversion.
    with tempfile.TemporaryDirectory(prefix="docmind_smoke_profile_") as profile:
        for source, target in (("docx", "doc:MS Word 97"), ("pptx", "ppt:MS PowerPoint 97")):
            subprocess.run(
                [libreoffice_executable(), f"-env:UserInstallation={Path(profile).as_uri()}",
                 "--headless", "--convert-to", target, "--outdir", str(directory),
                 str(directory / f"sample.{source}")],
                check=True, timeout=120, capture_output=True,
            )
            assert (directory / f"sample.{target.split(':')[0]}").is_file()
    for file in sorted(directory.glob("sample.*")):
        documents = load_documents(str(directory), input_files=[str(file)])
        assert documents and marker in "\n".join(doc.text for doc in documents), file.name
        print(f"PASS {file.suffix} extraction on read-only container", flush=True)

    embedding = MockEmbedding(embed_dim=8)
    with request_state(state):
        index = create_index(documents, embed_model=embedding, chunk_size=256, chunk_overlap=32)
    cache = session_directory("cache", state)
    index.storage_context.persist(persist_dir=str(cache))
    from llama_index.core import StorageContext, load_index_from_storage

    restored = load_index_from_storage(
        StorageContext.from_defaults(persist_dir=str(cache)), embed_model=embedding
    )
    assert restored.as_retriever(similarity_top_k=1).retrieve(marker)
    print("PASS tokenizer, indexing, persistence and retrieval without network", flush=True)
    with tempfile.NamedTemporaryFile(dir=Path.home() / ".cache") as handle:
        handle.write(b"cache permission check")
finally:
    clear_session_workspace(state)
    assert not directory.exists(), "Session workspace was not cleaned"
print("PASS container integration checks", flush=True)
