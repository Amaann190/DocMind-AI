"""Local React API. Run one worker; document state stays in this process."""
import asyncio
import hashlib
import io
import json
import secrets
import re
import tempfile
import threading
import time
from contextlib import contextmanager
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from fastapi import FastAPI, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field, model_validator
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.staticfiles import StaticFiles

from backend import history

from utils import helpers, llama_index as indexer, ollama
from utils.r2r import R2RClient, clear_remote_documents
from utils.runtime import request_state
from utils.session_workspace import clear_session_workspace, session_directory

app = FastAPI(title="DocMind Workspace", docs_url=None, redoc_url=None)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=["localhost", "127.0.0.1", "testserver"])
sessions = {}
sessions_lock = threading.Lock()
COOKIE = "docmind_workspace"


class Settings(BaseModel):
    provider: Literal["Ollama", "OpenAI", "LM Studio (Local AI)", "TabbyAPI"] = "Ollama"
    endpoint: str = "http://localhost:11434"
    chat_model: str = "llama3:latest"
    embedding_model: str = "nomic-embed-text:latest"
    api_key: str = ""
    style: Literal["Balanced", "Concise", "Detailed", "Bulleted", "Technical", "Simple"] = "Balanced"
    temperature: float = Field(default=.4, ge=0, le=1.5)
    top_k: int = Field(default=3, ge=1, le=10)
    chunk_size: int = Field(default=256, ge=64, le=8192)
    overlap: int = Field(default=32, ge=0, le=4096)
    similarity: float = Field(default=.3, ge=0, le=1)
    eco: bool = False
    r2r: bool = False
    r2r_url: str = "http://localhost:7272"
    r2r_key: str = ""

    @model_validator(mode="after")
    def validate_settings(self):
        for address in (self.endpoint, self.r2r_url):
            url = urlparse(address)
            if url.scheme not in ("http", "https") or not url.hostname or url.username or url.password or url.query or url.fragment:
                raise ValueError("Use a server URL without embedded credentials, query or fragment.")
        if self.overlap >= self.chunk_size:
            raise ValueError("Chunk overlap must be smaller than chunk size.")
        return self


def new_session():
    return {"config": Settings(), "messages": [], "sources": [], "warnings": [],
            "documents": [], "query_engine": None, "lock": threading.Lock(),
            "conversation_id": secrets.token_hex(16), "title": "New conversation", "renamed": False,
            "corpus_id": "empty", "source_inputs": {}, "cancel": threading.Event(),
            "touched": time.monotonic(), "status": "Ready", "progress": None}


def save_conversation(session):
    if not session.get("owner"):
        return
    if not session["renamed"]:
        session["title"] = next((m["content"][:70] for m in session["messages"] if m["role"] == "user"), "New conversation")
    history.save(session["owner"], session["conversation_id"], session["title"], {
        "messages": session["messages"], "renamed": session["renamed"], "corpus_id": session["corpus_id"],
        "rag_history_start": session.get("rag_history_start", 0), "chat_history_start": session.get("chat_history_start", 0)})


def open_conversation(session, record):
    payload = record["payload"]
    messages = payload["messages"]
    same_corpus = payload.get("corpus_id") == session["corpus_id"]
    session.update(conversation_id=record["id"], title=record["title"], messages=messages,
        renamed=payload.get("renamed", False),
        rag_history_start=payload.get("rag_history_start", 0) if same_corpus else len(messages),
        chat_history_start=payload.get("chat_history_start", 0) if same_corpus else len(messages))


@app.middleware("http")
async def local_boundary(request, call_next):
    if request.url.path.startswith("/api"):
        origin = request.headers.get("origin")
        if origin and urlparse(origin).netloc != request.headers.get("host"):
            return Response("Cross-origin requests are not allowed", status_code=403)
        if request.method not in ("GET", "HEAD") and request.headers.get("x-docmind-client") != "react":
            return Response("Missing application request header", status_code=403)
        length = request.headers.get("content-length", "0")
        if not length.isdigit() or int(length) > helpers.MAX_TOTAL_UPLOAD_BYTES + 1024 * 1024:
            return Response("Upload exceeds 100 MB", status_code=413)
    return await call_next(request)


def get_session(request):
    with sessions_lock:
        session = sessions.get(request.cookies.get(COOKIE))
    if session is None:
        raise HTTPException(401, "Workspace expired. Reload to start a new session.")
    session["touched"] = time.monotonic()
    return session


@contextmanager
def operation(session):
    if not session["lock"].acquire(blocking=False):
        raise HTTPException(409, "A task is still running. Wait for it to finish.")
    try:
        with request_state(session):
            configure(session)
            yield
            save_conversation(session)
    finally:
        session["progress"] = None
        session["status"] = "Ready"
        session["lock"].release()


def configure(session):
    cfg = session["config"]
    session.update(llm_backend=cfg.provider, selected_model=cfg.chat_model,
        ollama_endpoint=cfg.endpoint, ollama_embedding_model=cfg.embedding_model,
        openai_base_url=cfg.endpoint, openai_model=cfg.chat_model,
        openai_embedding_model=cfg.embedding_model, openai_api_key=cfg.api_key,
        temperature=cfg.temperature, top_k=cfg.top_k, similarity_cutoff=cfg.similarity,
        eco_mode=cfg.eco, system_prompt=f"You are DocMind, a careful document assistant. Answer in a {cfg.style.lower()} style. Be factual and cite supplied evidence.")


def snapshot(session):
    return {"settings": session["config"].model_dump(exclude={"api_key", "r2r_key"}),
        "conversation_id": session["conversation_id"], "title": session["title"],
        "conversations": history.listing(session["owner"]) if session.get("owner") else [],
        "sources": session["sources"], "messages": session["messages"],
        "warnings": session["warnings"], "busy": session["lock"].locked(),
        "status": session["status"], "progress": session["progress"],
        "has_index": session.get("query_engine") is not None or bool(session.get("r2r_document_ids")),
        "formats": sorted(helpers.ALLOWED_UPLOAD_EXTENSIONS)}


@app.get("/api/health")
def health():
    return {"ok": True}


@app.post("/api/session")
def init(request: Request, response: Response):
    token = request.cookies.get(COOKIE)
    with sessions_lock:
        # Local workspaces expire after eight idle hours; never discard remote cleanup IDs.
        for key, state in list(sessions.items()):
            if time.monotonic() - state["touched"] > 28800 and not state["lock"].locked() and not state.get("r2r_document_ids"):
                clear_session_workspace(state)
                del sessions[key]
        if token not in sessions:
            if len(sessions) >= 64:
                raise HTTPException(503, "Too many workspaces. Restart the local server to release them.")
            owner = hashlib.sha256(token.encode()).hexdigest() if token and re.fullmatch(r"[\w-]{43}", token) else ""
            saved = history.listing(owner) if owner else []
            if not saved:
                token = secrets.token_urlsafe(32)
                owner = hashlib.sha256(token.encode()).hexdigest()
            sessions[token] = new_session()
            sessions[token]["owner"] = owner
            if saved:
                open_conversation(sessions[token], history.read(owner, saved[0]["id"]))
            else:
                save_conversation(sessions[token])
        session = sessions[token]
    response.set_cookie(COOKIE, token, httponly=True, samesite="strict", max_age=365 * 86400)
    return snapshot(session)


@app.get("/api/state")
def state(request: Request):
    return snapshot(get_session(request))


def clear_corpus(session):
    if not clear_remote_documents(session):
        session["r2r_ingestion_failed"] = True
        raise HTTPException(502, "Remote documents could not be deleted. Restore the original R2R connection and retry.")
    session.update(query_engine=None, retriever=None, documents=[], sources=[], warnings=[],
        source_inputs={}, corpus_id="empty",
        last_doc_sources=[], last_doc_passages=[], last_rag_no_result=False,
        rag_history_start=len(session["messages"]), chat_history_start=len(session["messages"]))
    clear_session_workspace(session)


@app.put("/api/settings")
def settings(config: Settings, request: Request):
    session = get_session(request)
    with operation(session):
        old = session["config"]
        if not config.api_key and old.endpoint == config.endpoint and old.provider == config.provider:
            config.api_key = old.api_key
        if not config.r2r_key and old.r2r_url == config.r2r_url:
            config.r2r_key = old.r2r_key
        index_keys = ("provider", "endpoint", "embedding_model", "chunk_size", "overlap", "api_key", "r2r", "r2r_url", "r2r_key")
        if any(getattr(old, key) != getattr(config, key) for key in index_keys):
            clear_corpus(session)
        session["config"] = config
        configure(session)
    return snapshot(session)


@app.get("/api/models")
def models(request: Request):
    session = get_session(request)
    probe = {"config": session["config"]}
    # Discovery is read-only: do not lock out settings or chat while a server is slow.
    with request_state(probe):
        configure(probe)
        if probe["config"].provider == "Ollama":
            chat, embeddings = ollama.get_models(), ollama.get_embedding_models()
        else:
            chat = ollama.get_openai_models(probe["config"].endpoint, probe["config"].api_key)
            embeddings = chat
    return {"chat": chat, "embeddings": embeddings, "online": bool(chat)}


def index_documents(session, documents, kind, warnings=None):
    if not documents or len(documents) > 300 or sum(len(d.get_content()) for d in documents) > 4 * 1024 * 1024:
        raise ValueError("Use between 1 and 300 readable documents, with at most 4 MB of extracted text.")
    cfg = session["config"]
    session["status"] = "Connecting to embedding model"
    session_directory("cache", session)
    candidate = dict(session)
    def progress(done, total):
        session["status"] = "Building your document index"
        session["progress"] = {"done": done, "total": total}
    # Build separately: a failed replacement must not destroy the working index.
    with request_state(candidate):
        indexer.setup_embedding_model(cfg.embedding_model, cfg.chunk_size, cfg.overlap, cfg.provider, cfg.endpoint, cfg.api_key)
        candidate["llm"] = ollama.create_llm(cfg.chat_model, cfg.endpoint, cfg.api_key, backend=cfg.provider)
        indexer.create_query_engine(documents, progress_callback=progress)
    for key in ("embedding_model", "embedding_config", "llm", "retriever", "query_engine"):
        session[key] = candidate.get(key)
    session["documents"] = documents
    grouped = {}
    for doc in documents:
        name = doc.metadata.get("file_name") or doc.metadata.get("source") or "Document"
        source_id = doc.metadata.get("docmind_source_id", name)
        grouped.setdefault(source_id, {"name": name, "kind": doc.metadata.get("docmind_kind", kind), "texts": []})["texts"].append(doc.get_content())
    session["sources"] = [{"id": source_id, "name": group["name"], "kind": group["kind"], "status": "ready",
        "characters": sum(map(len, group["texts"])), "preview": "\n\n".join(group["texts"])[:16000]}
        for source_id, group in grouped.items()]
    session["warnings"] = warnings or []


def tag_documents(documents, source_id, kind):
    for doc in documents:
        doc.metadata.update(docmind_source_id=source_id, docmind_kind=kind)
        for key in ("docmind_source_id", "docmind_kind"):
            if key not in doc.excluded_embed_metadata_keys:
                doc.excluded_embed_metadata_keys.append(key)
            if key not in doc.excluded_llm_metadata_keys:
                doc.excluded_llm_metadata_keys.append(key)
    return documents


def corpus_changed(session):
    corpus_id = secrets.token_hex(16) if session.get("query_engine") or session.get("r2r_document_ids") else "empty"
    session.update(corpus_id=corpus_id, rag_history_start=len(session["messages"]),
                   chat_history_start=len(session["messages"]), last_doc_sources=[], last_doc_passages=[])


def commit_documents(session, documents, warnings=None):
    # ponytail: rebuild the bounded (300 documents / 4 MB text) index per edit;
    # use incremental indexing if measured edit latency becomes a problem.
    failed = [s for s in session["sources"] if s.get("status") == "error"]
    if documents:
        index_documents(session, documents, "file", warnings)
    else:
        session.update(documents=[], sources=[], query_engine=None, retriever=None, warnings=warnings or [])
    ready_ids = {s["id"] for s in session["sources"]}
    session["sources"].extend(s for s in failed if s["id"] not in ready_ids)
    corpus_changed(session)


def source_record(session, source_id):
    source = next((s for s in session["sources"] if s["id"] == source_id), None)
    if source is None:
        raise HTTPException(404, "Source not found in this workspace.")
    return source


def load_input(session, item, source_id):
    warnings = []
    with tempfile.TemporaryDirectory(dir=session_directory("data")) as directory:
        upload = io.BytesIO(item["data"])
        upload.name = item["name"]
        helpers.save_uploaded_file(upload, directory)
        docs = indexer.load_documents(directory, input_files=[str(Path(directory) / item["name"])], errors=warnings)
    if not docs:
        raise ValueError("; ".join(map(str, warnings)) or "No readable text found. Replace this file with a readable copy.")
    return tag_documents(docs, source_id, "file"), warnings


def remote_upload(session, item):
    cfg = session["config"]
    client = R2RClient(cfg.r2r_url, cfg.r2r_key)
    ids = []
    session["r2r_document_origin"] = {"base_url": client.base_url, "api_key": cfg.r2r_key}
    try:
        with tempfile.TemporaryDirectory(dir=session_directory("data")) as directory:
            upload = io.BytesIO(item["data"])
            upload.name = item["name"]
            helpers.save_uploaded_file(upload, directory)
            client.upload_documents([str(Path(directory) / item["name"])], ids)
        return ids[0]
    except Exception:
        # Preserve timed-out uploads for a later owned-ID cleanup.
        failed = client.delete_documents(ids)
        session.setdefault("r2r_document_ids", []).extend(failed)
        if failed:
            session["r2r_ingestion_failed"] = True
        raise


def put_source(session, source_id, item):
    previous = next((s for s in session["sources"] if s["id"] == source_id), None)
    retained = sum(len(v.get("data", b"")) for k, v in session["source_inputs"].items() if k != source_id)
    if retained + len(item.get("data", b"")) > helpers.MAX_TOTAL_UPLOAD_BYTES:
        raise ValueError("This workspace can retain up to 100 MB of files. Remove a source first.")
    if session["config"].r2r:
        remote_id = remote_upload(session, item)
        if previous and previous.get("remote_id"):
            client = R2RClient(**session["r2r_document_origin"])
            if client.delete_documents([previous["remote_id"]]):
                failed = client.delete_documents([remote_id])
                session.setdefault("r2r_document_ids", []).extend(failed)
                if failed:
                    session["r2r_ingestion_failed"] = True
                raise ValueError("The old remote source could not be removed. It has been kept; retry when R2R is available.")
            session["r2r_document_ids"].remove(previous["remote_id"])
        session.setdefault("r2r_document_ids", []).append(remote_id)
        session["sources"] = [s for s in session["sources"] if s["id"] != source_id] + [{
            "id": source_id, "name": item["name"], "kind": "r2r", "status": "ready", "remote_id": remote_id,
            "characters": 0, "preview": "This document is indexed on your R2R server."}]
        corpus_changed(session)
    else:
        docs, warnings = load_input(session, item, source_id)
        remaining = [d for d in session["documents"] if d.metadata.get("docmind_source_id") != source_id]
        commit_documents(session, remaining + docs, warnings)
    session["source_inputs"][source_id] = item


async def read_uploads(files):
    if not files or len(files) > 10:
        raise HTTPException(400, "Choose 1–10 files.")
    uploads, total = [], 0
    try:
        for file in files:
            helpers.safe_uploaded_filename(file.filename)
            data = await file.read(helpers.MAX_UPLOAD_FILE_BYTES + 1)
            total += len(data)
            if len(data) > helpers.MAX_UPLOAD_FILE_BYTES or total > helpers.MAX_TOTAL_UPLOAD_BYTES:
                raise ValueError("Use files under 25 MB each and 100 MB combined.")
            item = io.BytesIO(data)
            item.name = file.filename
            uploads.append(item)
        helpers.validate_uploaded_files(uploads)
        return uploads
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    finally:
        for file in files:
            await file.close()


def ingest_uploads(session, uploads):
    with operation(session):
        session["status"] = "Reading your documents"
        warnings = []
        for upload in uploads:
            if len(session["sources"]) >= 300:
                raise ValueError("Remove a source before adding more than 300 sources.")
            source_id = secrets.token_hex(16)
            item = {"kind": "file", "name": upload.name, "data": upload.getvalue()}
            retained = sum(len(v.get("data", b"")) for v in session["source_inputs"].values())
            if retained + len(item["data"]) > helpers.MAX_TOTAL_UPLOAD_BYTES:
                raise ValueError("This workspace can retain up to 100 MB of files. Remove a source first.")
            try:
                put_source(session, source_id, item)
                warnings.extend(session["warnings"])
            except Exception as error:
                session["source_inputs"][source_id] = item
                session["sources"].append({"id": source_id, "name": item["name"], "kind": "file", "status": "error",
                    "characters": 0, "preview": "", "error": str(error)})
                warnings.append(f"{item['name']}: {error}")
        session["warnings"] = warnings
    return snapshot(session)


@app.post("/api/upload")
async def upload(request: Request, files: list[UploadFile]):
    session = get_session(request)
    uploads = await read_uploads(files)
    try:
        return await asyncio.to_thread(ingest_uploads, session, uploads)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(400, str(error)) from error


class ImportSource(BaseModel):
    kind: Literal["website", "github", "sample"]
    value: str = Field(default="", max_length=4000)


@app.post("/api/import")
def import_source(source: ImportSource, request: Request):
    session = get_session(request)
    try:
        with operation(session):
            if session["config"].r2r:
                raise ValueError("R2R supports file uploads. Turn it off to import this source.")
            session["status"] = "Reading source content"
            warnings = []
            if source.kind == "website":
                docs = helpers.load_website_documents(source.value.splitlines())
            elif source.kind == "github":
                repo = helpers.normalize_github_repo(source.value)
                directory = helpers.clone_github_repo(repo)
                if not directory:
                    raise ValueError("Repository could not be cloned. Check its name and public access.")
                try:
                    docs = indexer.load_documents(directory, errors=warnings)
                finally:
                    helpers.remove_dir_retry(directory)
            else:
                from llama_index.core.schema import Document
                docs = [Document(text=("Northstar Studio — Project brief (sample document).\n"
                    "The Aurora website redesign launches on November 12, 2026. The budget is $24,000. "
                    "Maya Chen leads design and Omar Shah leads engineering. The project covers a new "
                    "homepage, product pages, and accessible navigation. The first design review is October 22. "
                    "Success means a 20 percent improvement in sign-up conversion and a page load time under 2 seconds. "
                    "The launch depends on content approval by November 5. The project does not include a mobile app."),
                    metadata={"file_name": "Northstar — Project brief.txt"})]
            grouped = {}
            for doc in docs:
                name = doc.metadata.get("file_name") or doc.metadata.get("source") or "Document"
                grouped.setdefault(name, []).append(doc)
            for group in grouped.values():
                tag_documents(group, secrets.token_hex(16), source.kind)
            commit_documents(session, session["documents"] + docs, warnings)
        return snapshot(session)
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(400, str(error)) from error


@app.delete("/api/sources")
def delete_sources(request: Request):
    session = get_session(request)
    with operation(session):
        clear_corpus(session)
    return snapshot(session)


@app.delete("/api/sources/{source_id}")
def delete_source(source_id: str, request: Request):
    session = get_session(request)
    with operation(session):
        source = source_record(session, source_id)
        if source.get("remote_id"):
            if R2RClient(**session["r2r_document_origin"]).delete_documents([source["remote_id"]]):
                raise HTTPException(502, "Remote source could not be deleted. Retry when R2R is available.")
            session["r2r_document_ids"].remove(source["remote_id"])
        elif source.get("status") != "error":
            docs = [d for d in session["documents"] if d.metadata.get("docmind_source_id") != source_id]
            try:
                commit_documents(session, docs)
            except Exception as error:
                raise HTTPException(400, str(error)) from error
        session["sources"] = [s for s in session["sources"] if s["id"] != source_id]
        session["source_inputs"].pop(source_id, None)
        corpus_changed(session)
    return snapshot(session)


@app.post("/api/sources/{source_id}/retry")
def retry_source(source_id: str, request: Request):
    session = get_session(request)
    with operation(session):
        source_record(session, source_id)
        try:
            item = session["source_inputs"].get(source_id)
            if item:
                put_source(session, source_id, item)
            else:
                # Imported website/repository text can be reindexed without another network fetch.
                commit_documents(session, session["documents"])
        except Exception as error:
            raise HTTPException(400, str(error)) from error
    return snapshot(session)


@app.put("/api/sources/{source_id}")
async def replace_source(source_id: str, request: Request, files: list[UploadFile]):
    session = get_session(request)
    uploads = await read_uploads(files)
    if len(uploads) != 1:
        raise HTTPException(400, "Choose exactly one replacement file.")
    def replace():
        with operation(session):
            source_record(session, source_id)
            try:
                put_source(session, source_id, {"kind": "file", "name": uploads[0].name, "data": uploads[0].getvalue()})
            except Exception as error:
                raise HTTPException(400, str(error)) from error
        return snapshot(session)
    return await asyncio.to_thread(replace)


@app.post("/api/conversations")
def new_conversation(request: Request):
    session = get_session(request)
    with operation(session):
        save_conversation(session)
        session.update(conversation_id=secrets.token_hex(16), title="New conversation", renamed=False,
                       messages=[], rag_history_start=0, chat_history_start=0)
    return snapshot(session)


@app.get("/api/conversations")
def search_conversations(request: Request, q: str = ""):
    return history.listing(get_session(request)["owner"], q[:200])


class ConversationTitle(BaseModel):
    title: str = Field(min_length=1, max_length=100)


@app.put("/api/conversations/{conversation_id}")
def rename_conversation(conversation_id: str, title: ConversationTitle, request: Request):
    session = get_session(request)
    if not title.title.strip():
        raise HTTPException(400, "Enter a conversation name.")
    with operation(session):
        record = history.read(session["owner"], conversation_id)
        if not record:
            raise HTTPException(404, "Conversation not found.")
        record["payload"]["renamed"] = True
        history.save(session["owner"], conversation_id, title.title.strip(), record["payload"])
        if conversation_id == session["conversation_id"]:
            session.update(title=title.title.strip(), renamed=True)
    return snapshot(session)


@app.post("/api/conversations/{conversation_id}/open")
def switch_conversation(conversation_id: str, request: Request):
    session = get_session(request)
    with operation(session):
        record = history.read(session["owner"], conversation_id)
        if not record:
            raise HTTPException(404, "Conversation not found.")
        save_conversation(session)
        open_conversation(session, record)
    return snapshot(session)


@app.delete("/api/conversations/{conversation_id}")
def delete_conversation(conversation_id: str, request: Request):
    session = get_session(request)
    with operation(session):
        if not history.read(session["owner"], conversation_id):
            raise HTTPException(404, "Conversation not found.")
        history.delete(session["owner"], conversation_id)
        if conversation_id == session["conversation_id"]:
            session.update(conversation_id=secrets.token_hex(16), title="New conversation", renamed=False,
                           messages=[], rag_history_start=0, chat_history_start=0)
    return snapshot(session)


@app.delete("/api/chat")
def clear_chat(request: Request):
    session = get_session(request)
    with operation(session):
        session.update(messages=[], chat_history_start=0, rag_history_start=0)
    return snapshot(session)


@app.delete("/api/session")
def reset(request: Request):
    session = get_session(request)
    with operation(session):
        clear_corpus(session)
        history.delete(session["owner"])
        session.update(conversation_id=secrets.token_hex(16), title="New conversation", renamed=False)
        session.update(config=Settings(), messages=[], chat_history_start=0, rag_history_start=0)
    return snapshot(session)


class Question(BaseModel):
    text: str = Field(min_length=1, max_length=6000)
    without_documents: bool = False
    regenerate: bool = False


@app.post("/api/chat")
async def chat(question: Question, request: Request):
    session = get_session(request)
    # Acquire before sending stream headers, so concurrent requests get a real 409.
    if not session["lock"].acquire(False):
        raise HTTPException(409, "A task is already running.")
    if question.regenerate:
        messages = session["messages"]
        if len(messages) < 2 or messages[-1]["role"] != "assistant" or messages[-2]["role"] != "user":
            session["lock"].release()
            raise HTTPException(400, "There is no answer to regenerate.")
        question.text = messages[-2]["content"]
        question.without_documents = messages[-2].get("without_documents", False)
        session["messages"] = messages[:-2]
        for key in ("rag_history_start", "chat_history_start"):
            session[key] = min(session.get(key, 0), len(session["messages"]))
    session["cancel"].clear()
    queue = asyncio.Queue()
    loop = asyncio.get_running_loop()
    def emit(payload):
        if not loop.is_closed():
            loop.call_soon_threadsafe(queue.put_nowait, payload)
    def generate():
        content = ""
        chunks = None
        try:
            with request_state(session):
                configure(session)
                session["status"] = "Reading and thinking"
                session["last_doc_passages"] = []
                session["last_rag_no_result"] = False
                session["messages"].append({"role": "user", "content": question.text, "without_documents": question.without_documents})
                cfg = session["config"]
                if cfg.r2r and not question.without_documents:
                    if not session.get("r2r_document_ids") or session.get("r2r_ingestion_failed"):
                        raise ValueError("Upload documents to R2R before asking a question.")
                    chunks = [R2RClient(cfg.r2r_url, cfg.r2r_key).rag(question.text, session["r2r_document_ids"])]
                elif session.get("query_engine") and not question.without_documents:
                    chunks = ollama.context_chat(question.text, session["query_engine"])
                else:
                    chunks = ollama.chat(question.text)
                for chunk in chunks:
                    if session["cancel"].is_set():
                        break
                    content += chunk
                    emit({"type": "token", "text": chunk})
                sources = session.get("last_doc_passages", [])
                message = {"role": "assistant", "content": content, "sources": sources,
                    "stopped": session["cancel"].is_set(),
                    "no_evidence": session.get("last_rag_no_result", False)}
                session["messages"].append(message)
                emit({"type": "done", "message": message})
        except Exception as error:
            message = {"role": "assistant", "content": content + "\n\nResponse interrupted: " + str(error), "sources": []}
            session["messages"].append(message)
            emit({"type": "error", "message": message})
        finally:
            try:
                if hasattr(chunks, "close"):
                    chunks.close()
                save_conversation(session)
            except Exception:
                emit({"type": "warning", "text": "The answer is available, but could not be saved to disk. Export it before closing."})
            finally:
                session["status"] = "Ready"
                session["lock"].release()
                emit(None)
    threading.Thread(target=generate, daemon=True).start()
    async def events():
        try:
            while True:
                value = await queue.get()
                if value is None:
                    break
                yield json.dumps(value) + "\n"
        finally:
            if session["lock"].locked():
                session["cancel"].set()
    return StreamingResponse(events(), media_type="application/x-ndjson", headers={"Cache-Control": "no-store"})


@app.post("/api/chat/stop")
def stop_chat(request: Request):
    session = get_session(request)
    session["cancel"].set()
    return {"stopping": True}


@app.get("/api/export")
def export(request: Request):
    from docx import Document
    session = get_session(request)
    document = Document()
    document.add_heading("DocMind — Conversation", 0)
    for message in list(session["messages"]):
        document.add_heading(message["role"].title(), 2)
        document.add_paragraph(message["content"])
        for source in message.get("sources", []):
            document.add_paragraph(f"[{source['number']}] {source['name']}\n{source['text']}")
    output = io.BytesIO()
    document.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="DocMind-conversation.docx"'})


dist = Path(__file__).resolve().parents[1] / "frontend" / "dist"
if dist.exists():
    app.mount("/assets", StaticFiles(directory=dist / "assets"), name="assets")

    @app.get("/{path:path}")
    def frontend(path: str):
        if path.startswith("api/"):
            raise HTTPException(404)
        return FileResponse(dist / "index.html")
