"""Real-format extraction checks. Requires the documented LibreOffice install."""

import json
import mailbox
import struct
import subprocess
import tempfile
import unittest
import zipfile
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from utils.document_readers import DocumentTextReader, libreoffice_executable
from utils.helpers import ALLOWED_UPLOAD_EXTENSIONS
from utils.llama_index import load_documents

MARKER = "ORBIT-7429"
SENTENCE = f"The launch approval code is {MARKER}."


def create_samples(folder):
    """Generate synthetic documents; no downloads, private data, or model calls."""
    text_files = {
        "txt": SENTENCE,
        "md": "# Launch plan\n\n" + SENTENCE,
        "markdown": "# Launch plan\n\n" + SENTENCE,
        "csv": f"product,approval_code\nEngine,{MARKER}\nBrake,NEBULA-8831\n",
        "tsv": f"product\tapproval_code\nEngine\t{MARKER}\nBrake\tNEBULA-8831\n",
        "json": json.dumps({"approval_code": MARKER}),
        "jsonl": json.dumps({"note": "Launch plan"}) + "\n" + json.dumps({"approval_code": MARKER}),
        "html": "<p>The launch approval code is ORBIT&#45;7429.</p><script>HIDDEN_SCRIPT</script>",
        "htm": "<p>The launch approval code is ORBIT&#45;7429.</p>",
        "xml": "<?xml version='1.0'?><launch><code>ORBIT&#45;7429</code></launch>",
        "rtf": r"{\rtf1\ansi The launch approval code is ORBIT\'2d7429.}",
        "ipynb": json.dumps({
            "nbformat": 4, "nbformat_minor": 5, "metadata": {},
            "cells": [{"cell_type": "code", "execution_count": 1, "metadata": {},
                       "source": ["raise RuntimeError('NEVER_EXECUTE_NOTEBOOK')"],
                       "outputs": [{"output_type": "stream", "name": "stdout", "text": [SENTENCE]}]}],
        }),
    }
    for extension, content in text_files.items():
        (folder / f"sample.{extension}").write_text(content, encoding="utf-8")

    message = EmailMessage()
    message["From"] = "sender@example.test"
    message["To"] = "recipient@example.test"
    message["Subject"] = "Launch plan"
    message.set_content(SENTENCE, cte="base64")
    (folder / "sample.eml").write_bytes(message.as_bytes())
    box = mailbox.mbox(folder / "sample.mbox")
    try:
        box.add(message)
    finally:
        box.close()
    message = EmailMessage()
    message.set_content(f"<p>{SENTENCE}</p>", subtype="html", cte="base64")
    message.make_related()
    (folder / "sample.mhtml").write_bytes(message.as_bytes())

    with zipfile.ZipFile(folder / "sample.odt", "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", (
            '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
            f'<office:body><office:text><text:p>{SENTENCE}</text:p></office:text></office:body>'
            '</office:document-content>'
        ))
        archive.writestr("META-INF/manifest.xml", (
            '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">'
            '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
            '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
            '</manifest:manifest>'
        ))

    from docx import Document
    document = Document()
    document.add_paragraph(SENTENCE)
    document.save(folder / "sample.docx")

    from pptx import Presentation
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = SENTENCE
    presentation.save(folder / "sample.pptx")

    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(["product", "approval_code"])
    workbook.active.append(["Engine", MARKER])
    workbook.create_sheet("Second sheet").append(["SECOND-SHEET-9284"])
    workbook.save(folder / "sample.xlsx")
    workbook.close()

    from ebooklib import epub
    book = epub.EpubBook()
    book.set_identifier("docmind-test")
    book.set_title("Launch plan")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Launch", file_name="launch.xhtml", lang="en")
    chapter.content = f"<html><body><p>{SENTENCE}</p></body></html>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    epub.write_epub(folder / "sample.epub", book)

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 750 Td ({SENTENCE}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = stream
    writer.write(folder / "sample.pdf")

    # Minimal Outlook message with real OLE streams and Unicode MAPI properties.
    from extract_msg.ole_writer import OleWriter
    ole = OleWriter()
    properties = bytes(32) + struct.pack("<IIII", 0x340D0003, 6, 0x40000, 0)
    for tag, value in ((0x001A001F, "IPM.Note"), (0x0037001F, "Launch plan"), (0x1000001F, SENTENCE)):
        data = (value + "\0").encode("utf-16-le")
        ole.addEntry(f"__substg1.0_{tag:08X}", data)
        properties += struct.pack("<IIII", tag, 6, len(data), 0)
    ole.addEntry("__properties_version1.0", properties)
    ole.write(str(folder / "sample.msg"))

    # Create actual legacy Office files, never modern files with renamed suffixes.
    with tempfile.TemporaryDirectory(prefix="docmind_test_profile_") as directory:
        for source, conversion in (("docx", "doc:MS Word 97"),
                                   ("pptx", "ppt:MS PowerPoint 97"),
                                   ("xlsx", "xls:MS Excel 97")):
            result = subprocess.run(
                [libreoffice_executable(), f"-env:UserInstallation={Path(directory).as_uri()}",
                 "--headless", "--convert-to", conversion, "--outdir", str(folder),
                 str(folder / f"sample.{source}")],
                capture_output=True, timeout=120, check=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            target = folder / ("sample." + conversion.split(":")[0])
            if not target.is_file():
                raise RuntimeError(f"Could not create {target.name}: {result.stdout!r} {result.stderr!r}")


class DocumentFormatTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix="docmind_formats_")
        cls.addClassCleanup(cls.directory.cleanup)
        cls.folder = Path(cls.directory.name)
        create_samples(cls.folder)

    def test_all_25_formats_through_upload_loader(self):
        paths = sorted(self.folder.glob("sample.*"))
        self.assertEqual({path.suffix for path in paths}, ALLOWED_UPLOAD_EXTENSIONS)
        self.assertEqual(len(paths), 25)
        for path in paths:
            with self.subTest(extension=path.suffix):
                documents = load_documents(str(self.folder), input_files=[str(path)])
                text = "\n".join(document.text for document in documents)
                self.assertIn(MARKER, text)
                self.assertNotIn("HIDDEN_SCRIPT", text)
                self.assertTrue(all(doc.metadata["file_name"] == path.name for doc in documents))
                if path.suffix in {".csv", ".tsv"}:
                    self.assertIn(f"approval_code is {MARKER}", text)
                if path.suffix in {".xls", ".xlsx"}:
                    self.assertIn("SECOND-SHEET-9284", text)

    def test_repository_directory_uses_same_readers(self):
        documents = load_documents(str(self.folder))
        found = {}
        for document in documents:
            extension = Path(document.metadata["file_name"]).suffix
            found[extension] = found.get(extension, "") + document.text
        self.assertEqual(set(found), ALLOWED_UPLOAD_EXTENSIONS)
        for extension, text in found.items():
            with self.subTest(extension=extension):
                self.assertIn(MARKER, text)


class ReaderSafetyTests(unittest.TestCase):
    def test_unicode_text_and_xml_entities(self):
        from defusedxml.common import EntitiesForbidden
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.txt"
            for encoding in ("utf-8-sig", "utf-16", "utf-32", "cp1252"):
                with self.subTest(encoding=encoding):
                    path.write_bytes("Caf\u00e9 ORBIT-7429".encode(encoding))
                    self.assertEqual(DocumentTextReader().load_data(path)[0].text, "Caf\u00e9 ORBIT-7429")
            path = path.with_suffix(".xml")
            path.write_text('<!DOCTYPE a [<!ENTITY x "expanded">]><a>&x;</a>')
            with self.assertRaises(EntitiesForbidden):
                DocumentTextReader().load_data(path)

    def test_odt_expansion_limit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "large.odt"
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("content.xml", "x" * 100)
            with patch("utils.document_readers.MAX_UPLOAD_FILE_BYTES", 50):
                with self.assertRaisesRegex(ValueError, "extraction size limit"):
                    DocumentTextReader().load_data(path)

    def test_missing_legacy_converter_has_setup_instructions(self):
        with patch("utils.document_readers.shutil.which", return_value=None), \
                patch("utils.document_readers.Path.is_file", return_value=False):
            with self.assertRaisesRegex(RuntimeError, "Install LibreOffice"):
                libreoffice_executable()


if __name__ == "__main__":
    unittest.main()
